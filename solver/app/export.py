"""シフト表の Excel（xlsx）出力（要件 3.5 / openpyxl）。

「シフト表」シート：スタッフ×日のグリッド（土日祝は色付け、店内掲示用）
「集計」シート：個人別の労働時間・出勤日数・土日出勤、概算人件費
"""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .models import Schedule

WEEKDAY_JP = ["月", "火", "水", "木", "金", "土", "日"]

_center = Alignment(horizontal="center", vertical="center")
_bold = Font(bold=True)
_thin = Border(*(Side(style="thin", color="CCCCCC"),) * 4)
_weekend_fill = PatternFill("solid", fgColor="FCE4EC")
_header_fill = PatternFill("solid", fgColor="E0F2F1")
_off_font = Font(color="AAAAAA")


def _short_label(label: str) -> str:
    """『早番 9:00-18:00』→『早』のような1文字ラベル。"""
    return label[0] if label else "?"


def build_workbook(schedule: Schedule) -> Workbook:
    wb = Workbook()
    label_of = {p.id: _short_label(p.label) for p in schedule.patterns}

    # ---------------- シフト表 ----------------
    ws = wb.active
    ws.title = "シフト表"

    ws.cell(row=1, column=1, value="スタッフ").font = _bold
    ws.cell(row=1, column=1).fill = _header_fill
    for j, d in enumerate(schedule.days, start=2):
        c1 = ws.cell(row=1, column=j, value=int(d.date[8:10]))
        c2 = ws.cell(row=2, column=j, value=WEEKDAY_JP[d.weekday])
        for c in (c1, c2):
            c.alignment = _center
            c.font = _bold
            c.fill = _weekend_fill if d.is_weekend else _header_fill
        ws.column_dimensions[c1.column_letter].width = 4.5
    ws.cell(row=2, column=1).fill = _header_fill
    ws.column_dimensions["A"].width = 14

    for i, s in enumerate(schedule.staff, start=3):
        ws.cell(row=i, column=1, value=s.name).border = _thin
        for j, d in enumerate(schedule.days, start=2):
            p = s.assignments.get(d.date)
            c = ws.cell(row=i, column=j, value=label_of.get(p, "?") if p else "休")
            c.alignment = _center
            c.border = _thin
            if d.is_weekend:
                c.fill = _weekend_fill
            if not p:
                c.font = _off_font

    # 不足行（BEST_EFFORT 時のみ）
    extra_rows = 0
    if any(d.shortage > 0 for d in schedule.days):
        r = len(schedule.staff) + 3
        c0 = ws.cell(row=r, column=1, value="不足")
        c0.font = Font(bold=True, color="C62828")
        for j, d in enumerate(schedule.days, start=2):
            if d.shortage > 0:
                c = ws.cell(row=r, column=j, value=-d.shortage)
                c.font = Font(bold=True, color="C62828")
                c.alignment = _center
        extra_rows = 1

    # 凡例
    legend_row = len(schedule.staff) + extra_rows + 4
    ws.cell(row=legend_row, column=1, value="凡例:").font = _bold
    for k, p in enumerate(schedule.patterns):
        ws.cell(row=legend_row + 1 + k, column=1,
                value=f"{_short_label(p.label)} = {p.label}（{p.hours}h）")

    # ---------------- 集計 ----------------
    ws2 = wb.create_sheet("集計")
    headers = ["スタッフ", "労働時間(h)", "出勤日数", "土日出勤"]
    for j, h in enumerate(headers, start=1):
        c = ws2.cell(row=1, column=j, value=h)
        c.font = _bold
        c.fill = _header_fill
        ws2.column_dimensions[c.column_letter].width = 14

    for i, s in enumerate(schedule.staff, start=2):
        work_days = sum(1 for v in s.assignments.values() if v)
        ws2.cell(row=i, column=1, value=s.name)
        ws2.cell(row=i, column=2, value=s.total_hours)
        ws2.cell(row=i, column=3, value=work_days)
        ws2.cell(row=i, column=4, value=s.weekend_shifts)

    info_row = len(schedule.staff) + 3
    ws2.cell(row=info_row, column=1, value="求解ステータス").font = _bold
    ws2.cell(row=info_row, column=2, value=schedule.meta.status)
    if schedule.meta.estimated_cost_yen is not None:
        ws2.cell(row=info_row + 1, column=1, value="概算人件費(円)").font = _bold
        ws2.cell(row=info_row + 1, column=2, value=schedule.meta.estimated_cost_yen)

    return wb


def to_xlsx_bytes(schedule: Schedule) -> bytes:
    buf = BytesIO()
    build_workbook(schedule).save(buf)
    return buf.getvalue()
